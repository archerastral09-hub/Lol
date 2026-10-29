#!/usr/bin/env python3
"""
Récupère l'équipe affichée sur la fiche de chaque agence L'Adresse
(colonne « Site internet » de agences_ladresse.xlsx) et remplit la feuille « Contacts ».

- 1 ligne par membre de l'équipe (nom, prénom, fonction)
- ajoute le représentant légal des mentions légales s'il n'est pas dans l'équipe
- colonne « Représentant légal » = Oui / Non pour trier plus tard
- reprise automatique : les agences déjà présentes dans Contacts sont ignorées
- sauvegarde régulière + compte rendu à la fin

Usage :
    python extraire_equipes.py
    python extraire_equipes.py --fichier agences_ladresse.xlsx --delai 1.0 --limite 10
    python extraire_equipes.py --refaire        # repart de zéro
"""
import argparse
import re
import sys
import time
import unicodedata
from collections import Counter

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

FEUILLE_AGENCES = "Agences"
FEUILLE_CONTACTS = "Contacts"
ENTETES_BASE = [
    "ID agence", "SIREN", "Nom", "Prénom", "Fonction",
    "Téléphone professionnel", "Email professionnel", "LinkedIn", "Source",
]
ENTETES_EXTRA = ["Nom agence", "Représentant légal"]
ENTETES = ENTETES_BASE + ENTETES_EXTRA

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "fr-FR,fr;q=0.9",
}


def log(msg):
    print(msg, flush=True)


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def normaliser(txt):
    txt = unicodedata.normalize("NFD", txt or "")
    txt = "".join(c for c in txt if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9 ]", " ", txt.lower()).split()


def cle_nom(txt):
    """Clé indépendante de l'ordre NOM Prénom / Prénom NOM et des accents."""
    return " ".join(sorted(normaliser(txt)))


def est_majuscule(token):
    lettres = [c for c in token if c.isalpha()]
    return bool(lettres) and all(c.isupper() for c in lettres)


def separer_nom_prenom(complet):
    """'Rémy BONNEMBERGER' -> ('BONNEMBERGER', 'Rémy')."""
    tokens = complet.split()
    if not tokens:
        return "", ""
    if len(tokens) == 1:
        return tokens[0], ""
    i = next((k for k, t in enumerate(tokens) if est_majuscule(t)), None)
    if i is None:
        return " ".join(tokens[1:]), tokens[0]
    if i > 0:
        return " ".join(tokens[i:]), " ".join(tokens[:i])
    # commence par des majuscules : « NOM Prénom » ou tout en majuscules
    j = next((k for k, t in enumerate(tokens) if not est_majuscule(t)), None)
    if j is not None:
        return " ".join(tokens[:j]), " ".join(tokens[j:])
    return " ".join(tokens[1:]), tokens[0]


def extraire_representant_legal(soup):
    for label in soup.find_all("label"):
        txt = " ".join(normaliser(label.get_text()))
        if txt.startswith("representant legal"):
            span = label.find_next_sibling("span") or label.find_next("span")
            if span:
                return span.get_text(" ", strip=True)
    return ""


def extraire_equipe(html):
    soup = BeautifulSoup(html, "html.parser")
    equipe, vus = [], set()
    for bloc in soup.select("div.agent"):
        nom_tag = bloc.select_one(".agent-nom")
        if not nom_tag:
            continue
        nom_complet = nom_tag.get_text(" ", strip=True)
        if not nom_complet or "rejoindre" in nom_complet.lower():
            continue
        fonction_tag = bloc.select_one(".agent-fonction")
        fonction = fonction_tag.get_text(" ", strip=True) if fonction_tag else ""
        cle = (cle_nom(nom_complet), fonction.lower())
        if cle in vus:
            continue
        vus.add(cle)
        equipe.append((nom_complet, fonction))
    return equipe, extraire_representant_legal(soup)


# --------------------------------------------------------------------------- #
# Réseau
# --------------------------------------------------------------------------- #
def telecharger(session, url, tentatives=4, timeout=30):
    derniere = ""
    for n in range(1, tentatives + 1):
        try:
            r = session.get(url, headers=HEADERS, timeout=timeout)
            if r.status_code == 200:
                r.encoding = r.encoding or "utf-8"
                return r.text, ""
            derniere = f"HTTP {r.status_code}"
            if r.status_code == 404:
                break
        except requests.RequestException as e:
            derniere = type(e).__name__
        time.sleep(2 * n)
    return None, derniere


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def preparer_feuille_contacts(wb, refaire):
    if FEUILLE_CONTACTS not in wb.sheetnames:
        ws = wb.create_sheet(FEUILLE_CONTACTS)
        ws.append(ENTETES)
        return ws
    ws = wb[FEUILLE_CONTACTS]
    if refaire:
        ws.delete_rows(2, ws.max_row)
    existants = [c.value for c in ws[1]]
    for extra in ENTETES_EXTRA:
        if extra not in existants:
            ws.cell(row=1, column=len(existants) + 1, value=extra)
            existants.append(extra)
    return ws


def lire_agences(ws):
    entetes = [str(c.value).strip() if c.value else "" for c in ws[1]]
    idx = {nom: i for i, nom in enumerate(entetes)}
    for col in ("ID agence", "Nom agence", "Site internet"):
        if col not in idx:
            sys.exit(f"Colonne « {col} » introuvable dans la feuille {FEUILLE_AGENCES}")
    agences = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx["ID agence"]]:
            continue
        agences.append({
            "id": str(row[idx["ID agence"]]).strip(),
            "siren": row[idx["SIREN"]] if "SIREN" in idx else "",
            "nom": row[idx["Nom agence"]] or "",
            "url": (row[idx["Site internet"]] or "").strip(),
        })
    return agences


def agences_deja_traitees(ws_contacts):
    entetes = [c.value for c in ws_contacts[1]]
    i = entetes.index("ID agence")
    return {str(r[i]).strip() for r in ws_contacts.iter_rows(min_row=2, values_only=True) if r[i]}


def ajouter_lignes(ws, agence, equipe, rep_legal):
    cle_rep = cle_nom(rep_legal) if rep_legal else ""
    cles_equipe = {cle_nom(n) for n, _ in equipe}
    lignes = []
    for nom_complet, fonction in equipe:
        nom, prenom = separer_nom_prenom(nom_complet)
        est_rep = "Oui" if cle_rep and cle_nom(nom_complet) == cle_rep else "Non"
        lignes.append((nom, prenom, fonction, est_rep))
    if rep_legal and cle_rep not in cles_equipe:
        nom, prenom = separer_nom_prenom(rep_legal)
        lignes.append((nom, prenom, "Représentant légal (mentions légales)", "Oui"))
    for nom, prenom, fonction, est_rep in lignes:
        ws.append([
            agence["id"], agence["siren"], nom, prenom, fonction,
            "", "", "", agence["url"], agence["nom"], est_rep,
        ])
    return len(lignes)


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    p = argparse.ArgumentParser()
    p.add_argument("--fichier", default="agences_ladresse.xlsx")
    p.add_argument("--delai", type=float, default=1.0, help="pause entre deux fiches (s)")
    p.add_argument("--limite", type=int, default=0, help="nombre max d'agences (0 = toutes)")
    p.add_argument("--refaire", action="store_true", help="vide la feuille Contacts avant")
    p.add_argument("--sauvegarde-tous", type=int, default=20)
    args = p.parse_args()

    wb = load_workbook(args.fichier)
    agences = lire_agences(wb[FEUILLE_AGENCES])
    ws_contacts = preparer_feuille_contacts(wb, args.refaire)
    deja = agences_deja_traitees(ws_contacts)

    a_faire = [a for a in agences if a["id"] not in deja]
    if args.limite:
        a_faire = a_faire[: args.limite]
    log(f"{len(agences)} agences dans le fichier, {len(deja)} déjà traitées, {len(a_faire)} à traiter")

    session = requests.Session()
    stats = Counter()
    sans_equipe, en_erreur = [], []
    total_contacts = 0

    for n, ag in enumerate(a_faire, 1):
        prefixe = f"[{n}/{len(a_faire)}] {ag['id']} {ag['nom']}"
        if not ag["url"].startswith("http"):
            log(f"{prefixe} -> pas d'URL, ignorée")
            stats["sans_url"] += 1
            en_erreur.append((ag["id"], ag["nom"], "pas d'URL"))
            continue

        html, erreur = telecharger(session, ag["url"])
        if html is None:
            log(f"{prefixe} -> ERREUR ({erreur})")
            stats["erreur"] += 1
            en_erreur.append((ag["id"], ag["nom"], erreur))
        else:
            equipe, rep = extraire_equipe(html)
            nb = ajouter_lignes(ws_contacts, ag, equipe, rep)
            total_contacts += nb
            if nb == 0:
                log(f"{prefixe} -> aucune personne trouvée")
                stats["sans_equipe"] += 1
                sans_equipe.append((ag["id"], ag["nom"]))
            else:
                log(f"{prefixe} -> {nb} contact(s) (représentant légal : {rep or 'n/c'})")
                stats["ok"] += 1

        if n % args.sauvegarde_tous == 0:
            wb.save(args.fichier)
            log(f"   (sauvegarde intermédiaire : {total_contacts} contacts)")
        time.sleep(args.delai)

    wb.save(args.fichier)

    log("\n========== COMPTE RENDU ==========")
    log(f"Agences traitées avec succès : {stats['ok']}")
    log(f"Agences sans équipe affichée : {stats['sans_equipe']}")
    log(f"Agences en erreur / sans URL : {stats['erreur'] + stats['sans_url']}")
    log(f"Contacts ajoutés             : {total_contacts}")
    if sans_equipe:
        log("\nSans équipe : " + ", ".join(f"{i} ({n})" for i, n in sans_equipe))
    if en_erreur:
        log("\nEn erreur :")
        for i, n, e in en_erreur:
            log(f"  - {i} {n} : {e}")
        log("\nRelance le script : seules les agences manquantes seront retraitées.")
    log(f"\nFichier mis à jour : {args.fichier} (feuille {FEUILLE_CONTACTS})")


if __name__ == "__main__":
    main()
